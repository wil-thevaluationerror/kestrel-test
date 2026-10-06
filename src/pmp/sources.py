"""Readable copies of the source documents, so a trace can be opened rather than described.

A lineage string tells a reviewer where a number came from. That is only useful if they can
get to the document without asking someone for the file. This module does two things:

1. Renders each ingested file into a cell grid keyed by the same reference the lineage uses,
   so the summary can show the source and highlight the exact cell behind a figure.
2. Copies the ingested files next to the report, so the run directory is a self-contained
   evidence pack that can be filed or handed to an auditor.

The copies are copies. Nothing here writes to data/raw/.
"""

from __future__ import annotations

import csv
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .hashing import sha256_file
from .paths import display_path

SOURCES_DIRNAME = "sources"
MAX_PREVIEW_ROWS = 200


@dataclass(frozen=True)
class Cell:
    ref: str | None  # None for header cells, which no lineage reference points at
    text: str


@dataclass(frozen=True)
class SheetPreview:
    name: str | None  # None for csv, which has no sheets
    columns: tuple[str, ...]
    rows: tuple[tuple[Cell, ...], ...]
    truncated: bool = False


@dataclass(frozen=True)
class SourceDocument:
    path: str  # repo-relative path of the original
    filename: str
    sha256: str
    kind: str  # workbook | delimited | text
    sheets: tuple[SheetPreview, ...] = ()
    text: str | None = None
    download: str | None = None  # path relative to summary.html, or None if not copied
    note: str | None = None


def _format_cell(value) -> str:
    """Render a cell the way the file shows it, without inventing precision."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _workbook_preview(path: Path) -> tuple[SheetPreview, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: list[SheetPreview] = []
        for worksheet in workbook.worksheets:
            raw_rows = list(worksheet.iter_rows(values_only=True))[:MAX_PREVIEW_ROWS]
            width = max((len(r) for r in raw_rows), default=0)
            rows = tuple(
                tuple(
                    Cell(
                        ref=f"{get_column_letter(column + 1)}{number}",
                        text=_format_cell(cells[column] if column < len(cells) else None),
                    )
                    for column in range(width)
                )
                for number, cells in enumerate(raw_rows, start=1)
            )
            sheets.append(
                SheetPreview(
                    name=worksheet.title,
                    columns=tuple(get_column_letter(i + 1) for i in range(width)),
                    rows=rows,
                    truncated=worksheet.max_row is not None and worksheet.max_row > MAX_PREVIEW_ROWS,
                )
            )
        return tuple(sheets)
    finally:
        workbook.close()


def _delimited_preview(path: Path) -> tuple[SheetPreview, ...]:
    """csv lineage references are `row N:<column>`, so the grid is keyed the same way."""
    with path.open(newline="", encoding="utf-8") as handle:
        records = list(csv.reader(handle))
    if not records:
        return (SheetPreview(name=None, columns=(), rows=()),)

    header, *data = records
    rows = [tuple(Cell(ref=None, text=value) for value in header)]
    for number, record in enumerate(data[:MAX_PREVIEW_ROWS], start=2):
        rows.append(
            tuple(
                Cell(ref=f"row {number}:{header[index]}" if index < len(header) else None, text=value)
                for index, value in enumerate(record)
            )
        )
    return (
        SheetPreview(
            name=None,
            columns=tuple(header),
            rows=tuple(rows),
            truncated=len(data) > MAX_PREVIEW_ROWS,
        ),
    )


def read_document(path: Path, repo_root: Path) -> SourceDocument:
    """Render one source file into something a reviewer can read on screen."""
    relative = display_path(path, repo_root)
    if path.suffix == ".xlsx":
        return SourceDocument(
            path=relative, filename=path.name, sha256=sha256_file(path),
            kind="workbook", sheets=_workbook_preview(path),
        )
    if path.suffix == ".csv":
        return SourceDocument(
            path=relative, filename=path.name, sha256=sha256_file(path),
            kind="delimited", sheets=_delimited_preview(path),
        )
    return SourceDocument(
        path=relative, filename=path.name, sha256=sha256_file(path),
        kind="text", text=path.read_text(encoding="utf-8"),
    )


@dataclass
class SourceBundle:
    """Every document the summary can open, keyed by repo-relative path."""

    documents: dict[str, SourceDocument] = field(default_factory=dict)

    def add(self, document: SourceDocument) -> None:
        self.documents[document.path] = document

    def get(self, relative_path: str) -> SourceDocument | None:
        return self.documents.get(relative_path)


def _unique_name(path: Path, taken: set[str]) -> str:
    """Flatten the raw directory tree into one folder without ever colliding."""
    candidate = path.name
    suffix = 2
    while candidate in taken:
        candidate = f"{path.stem}__{suffix}{path.suffix}"
        suffix += 1
    taken.add(candidate)
    return candidate


def collect(
    paths: list[Path], repo_root: Path, outputs_dir: Path, copy_files: bool
) -> SourceBundle:
    """Read every document and, unless asked not to, copy it next to the report.

    `copy_files` exists because shipping the underlying company reports alongside the summary
    is the right default for an internal evidence pack and the wrong default for something
    forwarded outside the fund. See docs/DATA_HANDLING.md.
    """
    bundle = SourceBundle()
    target_dir = outputs_dir / SOURCES_DIRNAME
    if copy_files:
        if target_dir.exists():
            shutil.rmtree(target_dir)
        target_dir.mkdir(parents=True, exist_ok=True)

    taken: set[str] = set()
    for path in sorted(set(paths)):
        document = read_document(path, repo_root)
        if copy_files:
            name = _unique_name(path, taken)
            shutil.copy2(path, target_dir / name)
            document = SourceDocument(
                **{**document.__dict__, "download": f"{SOURCES_DIRNAME}/{name}"}
            )
        else:
            document = SourceDocument(
                **{**document.__dict__, "note": "source copies omitted for this run"}
            )
        bundle.add(document)
    return bundle
