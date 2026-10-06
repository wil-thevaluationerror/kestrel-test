"""`make demo-break`: prove the pipeline refuses to report on data it cannot map.

A copy of the raw data is made in a scratch root and one source label is renamed there.
data/raw/ itself is never touched. The run should raise I01 and the gate should hold, which
is the whole argument of the project in one command.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

from .paths import Layout
from .xlsxio import save_workbook

DEFAULT_TARGET = ("RE-C", "2026Q2", "Rental Revenue", "Rent Roll Revenue")


def prepare(source: Layout, scratch_root: Path, target=DEFAULT_TARGET) -> tuple[Layout, Path]:
    """Copy config, templates, and raw data into a scratch root and rename one label there."""
    holding_id, period_label, old_label, new_label = target
    if scratch_root.exists():
        shutil.rmtree(scratch_root)
    scratch = Layout(root=scratch_root)
    shutil.copytree(source.config, scratch.config)
    shutil.copytree(source.templates, scratch.templates)
    shutil.copytree(source.raw, scratch.raw)

    matches = sorted(scratch.raw.glob(f"{period_label}/**/{holding_id}_*{period_label}.xlsx"))
    if not matches:
        raise FileNotFoundError(f"no {holding_id} {period_label} workbook to modify in {scratch.raw}")
    path = matches[0]

    workbook = load_workbook(path)
    renamed = False
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            if row and str(row[0].value).strip() == old_label:
                row[0].value = new_label
                renamed = True
    if not renamed:
        raise ValueError(f"label {old_label!r} not found in {path.name}")
    save_workbook(workbook, path)
    return scratch, path
