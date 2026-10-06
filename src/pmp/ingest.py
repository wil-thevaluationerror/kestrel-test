"""Read raw source reports, hash them, and stage them in one long format.

Staging keeps source values as text. Typing happens in normalize.py, so a letter in a
numeric cell surfaces as a located failure rather than a silent zero.

Every staged row is one source cell, which is what lets the summary trace an individual
number back to a file, sheet, and cell reference.

This module only ever opens files in data/raw/ for reading.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .contract import FUND_ENTITY, FundConfig
from .hashing import sha256_file
from .paths import Layout

HEADER_LABEL = "Line Item"
_FILENAME_PATTERN = re.compile(r"^(?P<entity>[A-Z]{2}-[A-Z]|FUND)_.*?_(?P<period>\d{4}Q\d)")


class IngestError(RuntimeError):
    """Raised when a source file cannot be read at all. Never swallowed: an unreadable
    submission is an operational event, not a missing value."""


@dataclass(frozen=True)
class StagedRow:
    """One source cell, with the address it came from."""

    record_type: str  # metric | cap_table | lp_capital | fund_summary
    entity_id: str  # holding_id, lp_id, or FUND
    holding_id: str  # FUND for fund-level files
    period_label: str
    source_label: str  # line-item label, holder name, or "" for row-keyed tables
    source_period: str  # "2026Q2" or "2026-05"
    field: str  # amount | shares | ownership_pct | called_usd | ...
    raw_value: str
    source_file: str
    source_sheet: str | None
    source_ref: str
    file_sha256: str


@dataclass(frozen=True)
class StagedReport:
    path: Path
    holding_id: str
    period_label: str
    file_sha256: str
    rows: tuple[StagedRow, ...]


@dataclass(frozen=True)
class DuplicateSubmission:
    """A file whose bytes were already ingested for this holding and period."""

    path: Path
    holding_id: str
    period_label: str
    file_sha256: str
    first_seen: str


# ------------------------------------------------------------------------- discovery


def discover(layout: Layout) -> list[Path]:
    """All raw submissions, sorted, so two runs see the same files in the same order."""
    return sorted(p for p in layout.raw.rglob("*") if p.suffix in {".xlsx", ".csv"})


def parse_submission_name(path: Path) -> tuple[str, str]:
    """Read holding and period from the filename submission standard.

    The convention is documented in docs/RUNBOOK.md. A file that does not follow it is
    rejected rather than guessed at.
    """
    match = _FILENAME_PATTERN.match(path.name)
    if not match:
        raise IngestError(
            f"{path.name}: does not follow <HOLDING_ID>_<slug>_<PERIOD>[_<suffix>].<ext>"
        )
    return match.group("entity"), match.group("period")


# --------------------------------------------------------------------------- parsers


def _split_at_header(sheet) -> tuple[list[tuple[int, list]], list[tuple[int, list]]]:
    """Split a sheet into its metadata block and its line-item block.

    Returns (metadata rows, line-item rows) as (row_number, cells) pairs. Metadata labels
    are staged too, so an unexpected label in the header block raises I01 instead of being
    quietly ignored.
    """
    metadata: list[tuple[int, list]] = []
    items: list[tuple[int, list]] = []
    seen_header = False
    for index, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not cells or cells[0] in (None, ""):
            continue
        if str(cells[0]).strip() == HEADER_LABEL:
            seen_header = True
            continue
        (items if seen_header else metadata).append((index, list(cells)))
    return metadata, items


def _cell_ref(row_number: int, column_index: int) -> str:
    return f"{get_column_letter(column_index)}{row_number}"


def _stage_label_value(
    *, holding_id: str, period_label: str, sheet_name: str, row_number: int,
    label: str, value, source_file: str, file_sha256: str, source_period: str,
    value_column: int,
) -> StagedRow:
    return StagedRow(
        record_type="metric",
        entity_id=holding_id,
        holding_id=holding_id,
        period_label=period_label,
        source_label=str(label).strip(),
        source_period=source_period,
        field="amount",
        raw_value="" if value is None else str(value),
        source_file=source_file,
        source_sheet=sheet_name,
        source_ref=_cell_ref(row_number, value_column),
        file_sha256=file_sha256,
    )


def _parse_two_column_xlsx(
    path: Path, holding_id: str, period_label: str, file_sha256: str, rel_path: str
) -> list[StagedRow]:
    """Property operating summaries and growth equity P&L / cash bridge sheets."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[StagedRow] = []
    try:
        for sheet in workbook.worksheets:
            metadata, items = _split_at_header(sheet)
            for row_number, cells in metadata + items:
                value = cells[1] if len(cells) > 1 else None
                rows.append(
                    _stage_label_value(
                        holding_id=holding_id, period_label=period_label,
                        sheet_name=sheet.title, row_number=row_number,
                        label=cells[0], value=value, source_file=rel_path,
                        file_sha256=file_sha256, source_period=period_label,
                        value_column=2,
                    )
                )
    finally:
        workbook.close()
    return rows


def _parse_venture_xlsx(
    path: Path, holding_id: str, period_label: str, file_sha256: str, rel_path: str
) -> list[StagedRow]:
    """Venture workbook: a Metrics sheet keyed by (label, period) and a Cap Table sheet."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[StagedRow] = []
    try:
        metrics = workbook["Metrics"]
        metadata, items = _split_at_header(metrics)
        for row_number, cells in metadata:
            rows.append(
                _stage_label_value(
                    holding_id=holding_id, period_label=period_label,
                    sheet_name=metrics.title, row_number=row_number,
                    label=cells[0], value=cells[1] if len(cells) > 1 else None,
                    source_file=rel_path, file_sha256=file_sha256,
                    source_period=period_label, value_column=2,
                )
            )
        for row_number, cells in items:
            rows.append(
                _stage_label_value(
                    holding_id=holding_id, period_label=period_label,
                    sheet_name=metrics.title, row_number=row_number,
                    label=cells[0], value=cells[2] if len(cells) > 2 else None,
                    source_file=rel_path, file_sha256=file_sha256,
                    source_period=str(cells[1]), value_column=3,
                )
            )
        rows.extend(
            _stage_cap_table_rows(
                records=[
                    {
                        "holder": str(cells[0]),
                        "share_class": str(cells[1]),
                        "shares": cells[2],
                        "ownership_pct": cells[3],
                    }
                    for cells in workbook["Cap Table"].iter_rows(min_row=2, values_only=True)
                    if cells and cells[0] not in (None, "")
                ],
                holding_id=holding_id, period_label=period_label,
                source_file=rel_path, sheet="Cap Table", file_sha256=file_sha256,
                first_row=2, ref_style="cell",
            )
        )
    finally:
        workbook.close()
    return rows


CAP_TABLE_FIELDS = ("share_class", "shares", "ownership_pct")
_CAP_TABLE_COLUMNS = {"share_class": 2, "shares": 3, "ownership_pct": 4}


def _stage_cap_table_rows(
    *, records: list[dict], holding_id: str, period_label: str, source_file: str,
    sheet: str | None, file_sha256: str, first_row: int, ref_style: str,
) -> list[StagedRow]:
    """One staged row per cap table cell, so a wrong ownership percentage points at a cell."""
    rows: list[StagedRow] = []
    for offset, record in enumerate(records):
        row_number = first_row + offset
        for field in CAP_TABLE_FIELDS:
            ref = (
                _cell_ref(row_number, _CAP_TABLE_COLUMNS[field])
                if ref_style == "cell"
                else f"row {row_number}:{field}"
            )
            rows.append(
                StagedRow(
                    record_type="cap_table",
                    entity_id=holding_id,
                    holding_id=holding_id,
                    period_label=period_label,
                    source_label=str(record["holder"]).strip(),
                    source_period=period_label,
                    field=field,
                    raw_value="" if record[field] is None else str(record[field]),
                    source_file=source_file,
                    source_sheet=sheet,
                    source_ref=ref,
                    file_sha256=file_sha256,
                )
            )
    return rows


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _parse_venture_metrics_csv(
    path: Path, holding_id: str, period_label: str, file_sha256: str, rel_path: str
) -> list[StagedRow]:
    return [
        StagedRow(
            record_type="metric",
            entity_id=holding_id,
            holding_id=holding_id,
            period_label=period_label,
            source_label=record["line_item"].strip(),
            source_period=record["period"].strip(),
            field="amount",
            raw_value=record["amount"].strip(),
            source_file=rel_path,
            source_sheet=None,
            source_ref=f"row {index}:amount",
            file_sha256=file_sha256,
        )
        for index, record in enumerate(_read_csv(path), start=2)
    ]


def _parse_venture_cap_table_csv(
    path: Path, holding_id: str, period_label: str, file_sha256: str, rel_path: str
) -> list[StagedRow]:
    return _stage_cap_table_rows(
        records=_read_csv(path),
        holding_id=holding_id,
        period_label=period_label,
        source_file=rel_path,
        sheet=None,
        file_sha256=file_sha256,
        first_row=2,
        ref_style="row",
    )


LP_FIELDS = ("commitment_usd", "called_usd", "distributed_usd")


def _parse_lp_capital_csv(
    path: Path, period_label: str, file_sha256: str, rel_path: str
) -> list[StagedRow]:
    rows: list[StagedRow] = []
    for index, record in enumerate(_read_csv(path), start=2):
        for field in LP_FIELDS:
            rows.append(
                StagedRow(
                    record_type="lp_capital",
                    entity_id=record["lp_id"].strip(),
                    holding_id=FUND_ENTITY,
                    period_label=period_label,
                    source_label=record["lp_id"].strip(),
                    source_period=period_label,
                    field=field,
                    raw_value=record[field].strip(),
                    source_file=rel_path,
                    source_sheet=None,
                    source_ref=f"row {index}:{field}",
                    file_sha256=file_sha256,
                )
            )
    return rows


FUND_SUMMARY_FIELDS = ("total_commitments_usd", "total_called_usd")


def _parse_fund_summary_csv(
    path: Path, period_label: str, file_sha256: str, rel_path: str
) -> list[StagedRow]:
    records = _read_csv(path)
    if len(records) != 1:
        raise IngestError(f"{rel_path}: expected exactly one fund summary row, got {len(records)}")
    return [
        StagedRow(
            record_type="fund_summary",
            entity_id=FUND_ENTITY,
            holding_id=FUND_ENTITY,
            period_label=period_label,
            source_label=FUND_ENTITY,
            source_period=period_label,
            field=field,
            raw_value=records[0][field].strip(),
            source_file=rel_path,
            source_sheet=None,
            source_ref=f"row 2:{field}",
            file_sha256=file_sha256,
        )
        for field in FUND_SUMMARY_FIELDS
    ]


# ------------------------------------------------------------------------- the stage


def _parse(path: Path, entity: str, period_label: str, config: FundConfig,
           file_sha256: str, rel_path: str) -> list[StagedRow]:
    """Dispatch to the parser for this submission's template and format."""
    if entity == FUND_ENTITY:
        if "lp-capital" in path.name:
            return _parse_lp_capital_csv(path, period_label, file_sha256, rel_path)
        if "fund-summary" in path.name:
            return _parse_fund_summary_csv(path, period_label, file_sha256, rel_path)
        raise IngestError(f"{rel_path}: unrecognised fund-level file")

    holding = config.holding(entity)
    if holding.template in {"real_estate", "growth_equity"}:
        return _parse_two_column_xlsx(path, entity, period_label, file_sha256, rel_path)
    if holding.template == "growth_venture":
        if path.suffix == ".xlsx":
            return _parse_venture_xlsx(path, entity, period_label, file_sha256, rel_path)
        if "cap-table" in path.name:
            return _parse_venture_cap_table_csv(path, entity, period_label, file_sha256, rel_path)
        return _parse_venture_metrics_csv(path, entity, period_label, file_sha256, rel_path)
    raise IngestError(f"{rel_path}: no parser for template {holding.template!r}")


def ingest(layout: Layout, config: FundConfig) -> tuple[list[StagedReport], list[DuplicateSubmission]]:
    """Hash and parse every raw submission.

    A file whose hash was already ingested for the same holding and period is reported as a
    duplicate and its rows are not staged, so one submission cannot be counted twice.
    """
    reports: list[StagedReport] = []
    duplicates: list[DuplicateSubmission] = []
    seen: dict[tuple[str, str, str], str] = {}

    for path in discover(layout):
        entity, period_label = parse_submission_name(path)
        digest = sha256_file(path)
        rel_path = str(path.relative_to(layout.root))
        key = (entity, period_label, digest)
        if key in seen:
            duplicates.append(
                DuplicateSubmission(
                    path=path, holding_id=entity, period_label=period_label,
                    file_sha256=digest, first_seen=seen[key],
                )
            )
            continue
        seen[key] = rel_path
        reports.append(
            StagedReport(
                path=path, holding_id=entity, period_label=period_label, file_sha256=digest,
                rows=tuple(_parse(path, entity, period_label, config, digest, rel_path)),
            )
        )
    return reports, duplicates
