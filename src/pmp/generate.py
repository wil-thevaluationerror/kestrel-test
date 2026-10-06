"""Seeded synthetic data generator.

Builds one fictional fund, nine holdings, two quarters of clean reports, then applies the
named defect transforms from config/fund.yaml on top. Clean data is built first and never
depends on the defect list, which is what makes `enabled: false` regenerate a clean file
exactly and the false-positive test meaningful.

Nothing here resembles a real fund, LP, or portfolio company.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, replace
from pathlib import Path
from random import Random

from .contract import FundConfig, Holding, Period
from .hashing import stable_seed
from .paths import Layout
from .textio import fmt_number, write_rows, write_text
from .xlsxio import save_workbook

from openpyxl import Workbook

THOUSANDS = 1000.0
LP_COMMITMENT_STEP = 250_000


@dataclass(frozen=True)
class LabelRow:
    """One line item as it will appear in a source report."""

    sheet: str
    label: str
    sub_period: str
    value: float
    unit: str  # USD or PCT


@dataclass(frozen=True)
class CapRow:
    holder: str
    share_class: str
    shares: float
    ownership_pct: float


@dataclass(frozen=True)
class HoldingPeriod:
    holding_id: str
    period_label: str
    labels: tuple[LabelRow, ...]
    cap_table: tuple[CapRow, ...] = ()


@dataclass(frozen=True)
class LPRow:
    lp_id: str
    commitment_usd: float
    called_usd: float
    distributed_usd: float


@dataclass(frozen=True)
class FundPeriod:
    period_label: str
    lps: tuple[LPRow, ...]
    total_commitments_usd: float
    total_called_usd: float


@dataclass(frozen=True)
class Universe:
    """Everything the writers need, after defects have been applied."""

    holdings: dict[tuple[str, str], HoldingPeriod]
    fund: dict[str, FundPeriod]
    duplicate_submissions: tuple[tuple[str, str], ...] = ()


# --------------------------------------------------------------------- clean universe


def _real_estate(holding: Holding, periods: list[Period], rng: Random) -> list[HoldingPeriod]:
    """Quarterly property operating summary. NOI is derived, not drawn, so the clean file
    is internally consistent."""
    revenue = round(rng.uniform(3_000_000, 5_000_000), -3)
    opex_ratio = rng.uniform(0.33, 0.40)
    occupancy = round(rng.uniform(88.0, 97.0), 1)
    debt = round(holding.invested_capital_usd * rng.uniform(1.3, 1.8), -3)

    out: list[HoldingPeriod] = []
    for index, period in enumerate(periods):
        if index > 0:
            revenue = round(revenue * rng.uniform(1.01, 1.06), -3)
            opex_ratio = opex_ratio * rng.uniform(0.98, 1.02)
            occupancy = round(min(99.0, max(85.0, occupancy + rng.uniform(-1.5, 1.5))), 1)
            debt = round(debt * 0.995, -3)
        opex = round(revenue * opex_ratio)
        noi = revenue - opex
        sheet = "Operating Summary"
        out.append(
            HoldingPeriod(
                holding_id=holding.holding_id,
                period_label=period.label,
                labels=(
                    LabelRow(sheet, "Rental Revenue", period.label, float(revenue), "USD"),
                    LabelRow(sheet, "Operating Expenses", period.label, float(opex), "USD"),
                    LabelRow(sheet, "Net Operating Income", period.label, float(noi), "USD"),
                    LabelRow(sheet, "Occupancy (%)", period.label, float(occupancy), "PCT"),
                    LabelRow(sheet, "Mortgage Balance", period.label, float(debt), "USD"),
                ),
            )
        )
    return out


def _growth_equity(holding: Holding, periods: list[Period], rng: Random) -> list[HoldingPeriod]:
    """Quarterly P&L plus a cash bridge that ties by construction."""
    revenue = round(rng.uniform(6_000_000, 14_000_000), -3)
    gross_margin = rng.uniform(0.58, 0.72)
    ebitda_margin = rng.uniform(-0.08, 0.12)
    cash = round(rng.uniform(8_000_000, 20_000_000), -3)
    debt = round(holding.invested_capital_usd * rng.uniform(0.2, 0.6), -3)

    out: list[HoldingPeriod] = []
    for index, period in enumerate(periods):
        if index > 0:
            revenue = round(revenue * rng.uniform(1.03, 1.12), -3)
            gross_margin = gross_margin * rng.uniform(0.99, 1.02)
            ebitda_margin = ebitda_margin + rng.uniform(-0.01, 0.02)
            debt = round(debt * 0.98, -3)
        gross_profit = round(revenue * gross_margin)
        ebitda = round(revenue * ebitda_margin)
        cash_begin = float(cash)
        cf_operating = float(round(ebitda * rng.uniform(0.6, 0.95)))
        cf_investing = float(-round(revenue * rng.uniform(0.03, 0.07)))
        cf_financing = float(round(rng.choice([0, 0, 1_000_000, -500_000])))
        cash_end = cash_begin + cf_operating + cf_investing + cf_financing
        cash = cash_end  # next quarter opens where this one closed

        out.append(
            HoldingPeriod(
                holding_id=holding.holding_id,
                period_label=period.label,
                labels=(
                    LabelRow("P&L", "Total Revenue", period.label, float(revenue), "USD"),
                    LabelRow("P&L", "Gross Profit", period.label, float(gross_profit), "USD"),
                    LabelRow("P&L", "EBITDA", period.label, float(ebitda), "USD"),
                    LabelRow("P&L", "Total Debt", period.label, float(debt), "USD"),
                    LabelRow("Cash Bridge", "Beginning Cash", period.label, cash_begin, "USD"),
                    LabelRow("Cash Bridge", "Cash from Operations", period.label, cf_operating, "USD"),
                    LabelRow("Cash Bridge", "Cash from Investing", period.label, cf_investing, "USD"),
                    LabelRow("Cash Bridge", "Cash from Financing", period.label, cf_financing, "USD"),
                    LabelRow("Cash Bridge", "Ending Cash", period.label, cash_end, "USD"),
                ),
            )
        )
    return out


def _cap_table(rng: Random) -> tuple[CapRow, ...]:
    """Ownership percentages are drawn so they sum to exactly 100.0 before any defect."""
    total_shares = 10_000_000
    founders = round(rng.uniform(38.0, 48.0), 1)
    seed_round = round(rng.uniform(14.0, 20.0), 1)
    series_a = round(rng.uniform(16.0, 22.0), 1)
    fund_stake = round(rng.uniform(8.0, 13.0), 1)
    option_pool = round(100.0 - founders - seed_round - series_a - fund_stake, 1)
    holders = (
        ("Founders", "Common", founders),
        ("Seed Investors", "Preferred Seed", seed_round),
        ("Series A Investors", "Preferred A", series_a),
        ("Lakeshore Partners Fund I", "Preferred A", fund_stake),
        ("Employee Option Pool", "Options", option_pool),
    )
    return tuple(
        CapRow(holder, share_class, float(round(total_shares * pct / 100.0)), float(pct))
        for holder, share_class, pct in holders
    )


def _growth_venture(holding: Holding, periods: list[Period], rng: Random) -> list[HoldingPeriod]:
    """Monthly revenue, quarterly burn and cash, plus a cap table."""
    monthly_revenue = round(rng.uniform(250_000, 600_000), -3)
    growth = rng.uniform(1.03, 1.08)
    cash = round(rng.uniform(6_000_000, 12_000_000), -3)

    out: list[HoldingPeriod] = []
    for period in periods:
        labels: list[LabelRow] = []
        for month in period.months:
            labels.append(LabelRow("Metrics", "Monthly Revenue", month, float(monthly_revenue), "USD"))
            monthly_revenue = round(monthly_revenue * growth, -3)
        net_burn = float(-round(rng.uniform(900_000, 1_800_000), -3))
        cash = float(round(cash + net_burn, -3))
        labels.append(LabelRow("Metrics", "Net Burn", period.label, net_burn, "USD"))
        labels.append(LabelRow("Metrics", "Ending Cash", period.label, cash, "USD"))
        out.append(
            HoldingPeriod(
                holding_id=holding.holding_id,
                period_label=period.label,
                labels=tuple(labels),
                cap_table=_cap_table(rng),
            )
        )
    return out


BUILDERS = {
    "real_estate": _real_estate,
    "growth_equity": _growth_equity,
    "growth_venture": _growth_venture,
}


def _lp_capital(config: FundConfig, rng: Random) -> dict[str, FundPeriod]:
    """40 fictional LPs whose commitments sum exactly to the fund total.

    Fund called capital is computed as the sum of the LP rows, so clean data ties by
    construction and R08 only fires when a defect breaks the tie.
    """
    lp_count = config.fund.lp_count
    weights = [rng.uniform(0.5, 2.5) for _ in range(lp_count)]
    scale = config.fund.total_commitments_usd / sum(weights)
    commitments = [
        round(w * scale / LP_COMMITMENT_STEP) * LP_COMMITMENT_STEP for w in weights
    ]
    commitments[0] += config.fund.total_commitments_usd - sum(commitments)

    called_pct = {"2026Q1": 0.62, "2026Q2": 0.68}
    distributed_pct = {"2026Q1": 0.11, "2026Q2": 0.15}

    fund: dict[str, FundPeriod] = {}
    for period in config.periods:
        rows = tuple(
            LPRow(
                lp_id=f"LP-{index + 1:03d}",
                commitment_usd=float(commitment),
                called_usd=float(round(commitment * called_pct[period.label])),
                distributed_usd=float(round(commitment * distributed_pct[period.label])),
            )
            for index, commitment in enumerate(commitments)
        )
        fund[period.label] = FundPeriod(
            period_label=period.label,
            lps=rows,
            total_commitments_usd=float(sum(r.commitment_usd for r in rows)),
            total_called_usd=float(sum(r.called_usd for r in rows)),
        )
    return fund


def build_clean_universe(config: FundConfig) -> Universe:
    """Clean reports for every holding and period. No defect logic runs here."""
    holdings: dict[tuple[str, str], HoldingPeriod] = {}
    for holding in config.holdings:
        rng = Random(stable_seed(config.seed, holding.holding_id))
        for record in BUILDERS[holding.template](holding, list(config.periods), rng):
            holdings[(record.holding_id, record.period_label)] = record
    fund = _lp_capital(config, Random(stable_seed(config.seed, "FUND")))
    return Universe(holdings=holdings, fund=fund)


# ---------------------------------------------------------------------------- defects


def _relabel(record: HoldingPeriod, old: str, new: str) -> HoldingPeriod:
    labels = tuple(
        replace(row, label=new) if row.label == old else row for row in record.labels
    )
    return replace(record, labels=labels)


def _shift_value(record: HoldingPeriod, label: str, delta: float) -> HoldingPeriod:
    labels = tuple(
        replace(row, value=row.value + delta) if row.label == label else row
        for row in record.labels
    )
    return replace(record, labels=labels)


def _rescale_usd(record: HoldingPeriod, divisor: float) -> HoldingPeriod:
    labels = tuple(
        replace(row, value=round(row.value / divisor, 1)) if row.unit == "USD" else row
        for row in record.labels
    )
    return replace(record, labels=labels)


def _shift_ownership(record: HoldingPeriod, holder: str, delta: float) -> HoldingPeriod:
    cap_table = tuple(
        replace(row, ownership_pct=round(row.ownership_pct + delta, 2))
        if row.holder == holder
        else row
        for row in record.cap_table
    )
    return replace(record, cap_table=cap_table)


def _shift_lp_called(period: FundPeriod, lp_id: str, delta: float) -> FundPeriod:
    """Only the LP row moves; the fund total stays as filed, which is what breaks the tie."""
    lps = tuple(
        replace(row, called_usd=row.called_usd + delta) if row.lp_id == lp_id else row
        for row in period.lps
    )
    return replace(period, lps=lps)


def apply_defects(universe: Universe, config: FundConfig) -> Universe:
    """Apply each enabled defect as a named transform. Order is fixed by config order."""
    holdings = dict(universe.holdings)
    fund = dict(universe.fund)
    duplicates: list[tuple[str, str]] = []

    for defect in config.defects:
        if not defect.enabled:
            continue
        key = (defect.holding_id, defect.period)
        if defect.id == "D01":
            holdings[key] = _shift_ownership(holdings[key], "Founders", -2.6)
        elif defect.id == "D02":
            holdings[key] = _relabel(holdings[key], "Total Revenue", "Net Sales")
        elif defect.id == "D04":
            holdings[key] = _shift_value(holdings[key], "Cash from Investing", 48_200.0)
        elif defect.id == "D07":
            duplicates.append(key)
        elif defect.id == "D08":
            holdings[key] = _rescale_usd(holdings[key], THOUSANDS)
        elif defect.id == "D10":
            fund[defect.period] = _shift_lp_called(fund[defect.period], "LP-017", -312_500.0)
        else:
            raise NotImplementedError(
                f"defect {defect.id} is enabled in config but has no transform; "
                "see docs/FAILURE_MODES.md for the deferred list"
            )

    return Universe(
        holdings=holdings, fund=fund, duplicate_submissions=tuple(duplicates)
    )


# ---------------------------------------------------------------------------- writers


def _slug(name: str) -> str:
    """Filesystem-safe holding slug with runs of separators collapsed."""
    cleaned = "".join(c.lower() if c.isalnum() else "-" for c in name)
    while "--" in cleaned:
        cleaned = cleaned.replace("--", "-")
    return cleaned.strip("-")


def _filename(holding_id: str, name: str, period_label: str, suffix: str, ext: str) -> str:
    """Submission naming standard: <HOLDING_ID>_<slug>_<PERIOD>[_<suffix>].<ext>.

    Ingestion reads holding and period from the filename, so the convention is part of the
    operating procedure (see docs/RUNBOOK.md), not a parsing convenience.
    """
    tail = f"_{suffix}" if suffix else ""
    return f"{holding_id}_{_slug(name)}_{period_label}{tail}.{ext}"


def _write_sheet_rows(worksheet, header: tuple[str, str], rows: list[tuple[str, float]]) -> None:
    worksheet.append(list(header))
    for label, value in rows:
        worksheet.append([label, value])


def _write_real_estate(path: Path, holding: Holding, period: Period, record: HoldingPeriod) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Operating Summary"
    sheet.append(["Property", holding.name])
    sheet.append(["Period Ending", period.period_end.isoformat()])
    sheet.append(["Reporting Currency", "USD"])
    sheet.append([])
    _write_sheet_rows(
        sheet, ("Line Item", "Amount"), [(r.label, r.value) for r in record.labels]
    )
    save_workbook(workbook, path)


def _write_growth_equity(path: Path, holding: Holding, period: Period, record: HoldingPeriod) -> None:
    workbook = Workbook()
    pnl = workbook.active
    pnl.title = "P&L"
    pnl.append(["Company", holding.name])
    pnl.append(["Period Ending", period.period_end.isoformat()])
    pnl.append([])
    _write_sheet_rows(
        pnl,
        ("Line Item", "Amount"),
        [(r.label, r.value) for r in record.labels if r.sheet == "P&L"],
    )

    bridge = workbook.create_sheet("Cash Bridge")
    bridge.append(["Company", holding.name])
    bridge.append(["Period Ending", period.period_end.isoformat()])
    bridge.append([])
    _write_sheet_rows(
        bridge,
        ("Line Item", "Amount"),
        [(r.label, r.value) for r in record.labels if r.sheet == "Cash Bridge"],
    )
    save_workbook(workbook, path)


def _venture_metric_rows(record: HoldingPeriod) -> list[list[str]]:
    return [
        [row.label, row.sub_period, fmt_number(row.value)] for row in record.labels
    ]


def _venture_cap_rows(record: HoldingPeriod) -> list[list[str]]:
    return [
        [row.holder, row.share_class, fmt_number(row.shares), fmt_number(row.ownership_pct)]
        for row in record.cap_table
    ]


def _write_venture_csv(base: Path, holding: Holding, period: Period, record: HoldingPeriod) -> None:
    write_rows(
        base.with_name(_filename(holding.holding_id, holding.name, period.label, "metrics", "csv")),
        ["line_item", "period", "amount"],
        _venture_metric_rows(record),
    )
    write_rows(
        base.with_name(_filename(holding.holding_id, holding.name, period.label, "cap-table", "csv")),
        ["holder", "share_class", "shares", "ownership_pct"],
        _venture_cap_rows(record),
    )


def _write_venture_xlsx(path: Path, holding: Holding, period: Period, record: HoldingPeriod) -> None:
    workbook = Workbook()
    metrics = workbook.active
    metrics.title = "Metrics"
    metrics.append(["Company", holding.name])
    metrics.append(["Period Ending", period.period_end.isoformat()])
    metrics.append([])
    metrics.append(["Line Item", "Period", "Amount"])
    for row in record.labels:
        metrics.append([row.label, row.sub_period, row.value])

    cap = workbook.create_sheet("Cap Table")
    cap.append(["Holder", "Share Class", "Shares", "Ownership (%)"])
    for row in record.cap_table:
        cap.append([row.holder, row.share_class, row.shares, row.ownership_pct])
    save_workbook(workbook, path)


def _write_fund_files(directory: Path, config: FundConfig, period: Period, record: FundPeriod) -> None:
    name = config.fund.name
    write_rows(
        directory / _filename("FUND", name, period.label, "lp-capital", "csv"),
        ["lp_id", "period_end", "commitment_usd", "called_usd", "distributed_usd"],
        [
            [
                row.lp_id,
                period.period_end.isoformat(),
                fmt_number(row.commitment_usd),
                fmt_number(row.called_usd),
                fmt_number(row.distributed_usd),
            ]
            for row in record.lps
        ],
    )
    write_rows(
        directory / _filename("FUND", name, period.label, "fund-summary", "csv"),
        ["period_end", "total_commitments_usd", "total_called_usd"],
        [
            [
                period.period_end.isoformat(),
                fmt_number(record.total_commitments_usd),
                fmt_number(record.total_called_usd),
            ]
        ],
    )


def write_universe(universe: Universe, config: FundConfig, layout: Layout) -> list[Path]:
    """Write every source report. Returns the files written, sorted for stable logs."""
    written: list[Path] = []
    for period in config.periods:
        for holding in config.holdings:
            record = universe.holdings[(holding.holding_id, period.label)]
            directory = layout.raw / period.label / holding.asset_class
            directory.mkdir(parents=True, exist_ok=True)
            if holding.template == "real_estate":
                path = directory / _filename(holding.holding_id, holding.name, period.label, "", "xlsx")
                _write_real_estate(path, holding, period, record)
                written.append(path)
            elif holding.template == "growth_equity":
                path = directory / _filename(holding.holding_id, holding.name, period.label, "", "xlsx")
                _write_growth_equity(path, holding, period, record)
                written.append(path)
            elif holding.report_format == "csv":
                base = directory / "placeholder.csv"
                _write_venture_csv(base, holding, period, record)
                written.extend(sorted(directory.glob(f"{holding.holding_id}_*{period.label}*.csv")))
            else:
                path = directory / _filename(holding.holding_id, holding.name, period.label, "", "xlsx")
                _write_venture_xlsx(path, holding, period, record)
                written.append(path)

        fund_dir = layout.raw / period.label / "fund"
        fund_dir.mkdir(parents=True, exist_ok=True)
        _write_fund_files(fund_dir, config, period, universe.fund[period.label])
        written.extend(sorted(fund_dir.glob(f"FUND_*{period.label}*.csv")))

    for holding_id, period_label in universe.duplicate_submissions:
        holding = config.holding(holding_id)
        period = config.period(period_label)
        directory = layout.raw / period_label / holding.asset_class
        original = directory / _filename(holding_id, holding.name, period_label, "", "xlsx")
        duplicate = directory / _filename(holding_id, holding.name, period_label, "resubmit", "xlsx")
        duplicate.write_bytes(original.read_bytes())
        written.append(duplicate)

    return sorted(set(written))


def write_ground_truth(config: FundConfig, layout: Layout) -> Path:
    """The defect manifest the tests score against. No timestamp: it must be reproducible."""
    payload = {
        "seed": config.seed,
        "fund": config.fund.name,
        "periods": [p.label for p in config.periods],
        "defects": [
            {
                "id": d.id,
                "description": d.description,
                "entity": d.holding_id,
                "period": d.period,
                "expected_rule": d.expected_rule,
            }
            for d in config.defects
            if d.enabled
        ],
        "clean_periods": [config.periods[0].label],
    }
    layout.ground_truth.parent.mkdir(parents=True, exist_ok=True)
    write_text(layout.ground_truth, json.dumps(payload, indent=2, sort_keys=True) + "\n")
    return layout.ground_truth


def generate(config: FundConfig, layout: Layout) -> list[Path]:
    """Full generation: clean universe, defect transforms, files, ground truth."""
    universe = apply_defects(build_clean_universe(config), config)
    files = write_universe(universe, config, layout)
    write_ground_truth(config, layout)
    return files
