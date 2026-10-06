"""I01 and I02, the two checks that run before any rule sees the data."""

import pytest
from openpyxl import load_workbook

from pmp.configio import load_config
from pmp.ingest import IngestError, ingest, parse_submission_name
from pmp.normalize import NormalizeError, normalize
from pmp.pipeline import run
from pmp.xlsxio import save_workbook


def _rename_label(layout, pattern: str, old: str, new: str):
    path = sorted(layout.raw.glob(pattern))[0]
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            if row and str(row[0].value).strip() == old:
                row[0].value = new
    save_workbook(workbook, path)
    return path


def test_i01_flags_an_unmapped_label_with_its_cell_address(mutable_project):
    # Arrange: rename a mapped label on an otherwise clean holding.
    _rename_label(mutable_project, "2026Q2/**/RE-C_*.xlsx", "Rental Revenue", "Rent Roll Revenue")

    result = run(mutable_project)

    unmapped = [e for e in result.queue.exceptions if e.rule_id == "I01" and e.holding_id == "RE-C"]
    assert len(unmapped) == 1
    assert unmapped[0].severity == "block"
    assert "Operating Summary!B6" in unmapped[0].evidence


def test_an_unmapped_value_never_enters_the_canonical_tables(mutable_project):
    """The value must be withheld rather than guessed at - that is why the gate exists."""
    _rename_label(mutable_project, "2026Q2/**/RE-C_*.xlsx", "Rental Revenue", "Rent Roll Revenue")

    result = run(mutable_project)

    revenue = [
        row for row in result.canonical.metrics
        if row.holding_id == "RE-C" and row.period_label == "2026Q2" and row.metric == "revenue"
    ]
    assert revenue == []


def test_i02_detects_the_duplicate_submission_in_the_seeded_data(project):
    config = load_config(project)
    reports, duplicates = ingest(project, config.fund)

    assert [(d.holding_id, d.period_label) for d in duplicates] == [("RE-A", "2026Q2")]
    assert all("resubmit" not in r.path.name for r in reports)


def test_a_duplicate_file_is_not_counted_twice(project, demo_resolutions):
    result = run(project, resolutions_path=demo_resolutions)

    revenue = [
        row for row in result.canonical.metrics
        if row.holding_id == "RE-A" and row.period_label == "2026Q2" and row.metric == "revenue"
    ]
    assert len(revenue) == 1


def test_a_file_outside_the_naming_standard_is_rejected_not_guessed(mutable_project):
    stray = mutable_project.raw / "2026Q2" / "real_estate" / "q2 numbers final.xlsx"
    stray.write_bytes(b"not a workbook")

    with pytest.raises(IngestError, match="does not follow"):
        run(mutable_project)


def test_parse_submission_name_reads_holding_and_period():
    from pathlib import Path

    assert parse_submission_name(Path("RE-A_harborview_2026Q2_resubmit.xlsx")) == ("RE-A", "2026Q2")


def test_a_non_numeric_value_in_a_mapped_cell_stops_the_run(mutable_project):
    """Never coerce. A letter where a number belongs is an operational event."""
    path = sorted(mutable_project.raw.glob("2026Q2/**/RE-C_*.xlsx"))[0]
    workbook = load_workbook(path)
    sheet = workbook["Operating Summary"]
    for row in sheet.iter_rows():
        if row and str(row[0].value).strip() == "Rental Revenue":
            row[1].value = "see note"
    save_workbook(workbook, path)

    config = load_config(mutable_project)
    reports, _ = ingest(mutable_project, config.fund)
    with pytest.raises(NormalizeError, match="cannot read 'see note' as a number"):
        normalize(reports, config, "run-test")
